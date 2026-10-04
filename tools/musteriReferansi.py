"""Yapay girdileri yerel BUP Yönetim ile çalıştıran bağımsız başvuru üreticisi.

Kullanım: BUP_Yonetim/.venv/bin/python tools/musteriReferansi.py KAYNAK_KLASOR CIKTI_JSON
Gerçek veri/config okunmaz. Geçici Excel'ler TemporaryDirectory içinde kalır.
"""

import hashlib
import json
import random
import sys
import tempfile
from pathlib import Path


def main():
    root = Path(sys.argv[1]).resolve()
    sys.path.insert(0, str(root))
    from openpyxl import Workbook, load_workbook
    from application.customer_tracking.facade import CustomerTrackingFacade
    from domain.musteri_takip.models import CustomerDataError
    from infrastructure.excel.customer_list_reader import CustomerListReader
    from infrastructure.export.customer_comparison_exporter import CustomerComparisonExcelExporter
    from infrastructure.export.visible_table_exporter import OpenpyxlVisibleTableExporter

    class Drivers:
        def __init__(self, values):
            self.values = values

        def drivers(self):
            return self.values

    class NoImage:
        def export(self, *args, **kwargs):
            raise AssertionError("Bu üretici resim veya gerçek veri okumaz.")

    def rows(names, depo=None, header=3):
        result = [["Yapay müşteri raporu"]]
        if depo:
            result.append([f"Cari Kategori 3 [TEST] {depo}"])
        while len(result) < header - 1:
            result.append([])
        result.append(["Kod", "Cari Ünvan", "Bakiye"])
        result.extend([[None, name, 0] for name in names])
        return result

    specs = [
        {"ad": "tekrar-ve-iki-yon", "eski": rows(["  Yapay Alfa  ", "Yapay Beta", "yapay alfa", None, "", "Yapay Gamma"]), "yeni": rows(["YAPAY ALFA", "Yapay Delta", "Yapay Delta"]), "plasiyerler": {}},
        {"ad": "arac-baslik", "eski": rows(["Yapay Alfa", "Yapay Beta"], "İZMİR ARAÇ 06", 15), "yeni": rows(["yapay alfa", "Yapay Yeni"]), "plasiyerler": {"06": "Yapay Plasiyer"}},
        {"ad": "unicode-ve-ic-bosluk", "eski": rows(["ışık", "IŞIK", "İşık", "Straße", "STRASSE", "Yapay  İki", "Yapay İki", "\u0085Yapay Kenar\u0085", "\ufeffYapay BOM\ufeff"]), "yeni": rows(["ışık", "strasse", "Yapay İki", "Yapay Kenar", "Yapay BOM"]), "plasiyerler": {}},
        {"ad": "bos-liste", "eski": rows([]), "yeni": rows([]), "plasiyerler": {}},
        {"ad": "sayisal-ve-bool", "eski": rows([123, True, False, 12.5]), "yeni": rows(["123", "True"]), "plasiyerler": {}},
        {"ad": "formul-metni", "eski": rows(["=YAPAY()", "+Yapay", "-Yapay", "@Yapay", "Yapay Normal"], "=YAPAY_BASLIK"), "yeni": rows([]), "plasiyerler": {}},
        {"ad": "tek-haneli-arac", "eski": rows(["Yapay B"], "Araç 7 / Yapay"), "yeni": rows([]), "plasiyerler": {"07": "Yapay / Plasiyer"}},
        {"ad": "unicode-arac-siniri", "eski": rows(["Yapay"], "Yapayç06"), "yeni": rows([]), "plasiyerler": {"06": "Yapay"}},
    ]
    rng = random.Random(1704)
    pool = ["Yapay Alfa", "yapay alfa", "İYAPAY", "ıyapay", "Yapay  A", "Yapay A", "ßYapay", "SSYapay", "", None, "  Yapay Kenar "]
    for index in range(24):
        specs.append({"ad": f"yapay-{index:02d}", "eski": rows(rng.choices(pool, k=rng.randrange(0, 18))), "yeni": rows(rng.choices(pool, k=rng.randrange(0, 18))), "plasiyerler": {}})

    def write_book(path, values):
        book = Workbook()
        for row in values:
            book.active.append(row)
        # Formül benzeri müşteri adlarını gerçek formül olarak kaydetme.
        for row in book.active:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        book.save(path)
        book.close()

    def cells(path):
        book = load_workbook(path, data_only=False)
        result = [{"ad": sheet.title, "satirlar": [list(row) for row in sheet.values], "birlesimler": sorted(str(r) for r in sheet.merged_cells.ranges)} for sheet in book]
        book.close()
        return result

    cases = []
    with tempfile.TemporaryDirectory(prefix="cal-musteri-reference-") as temp:
        folder = Path(temp)
        for spec in specs:
            old, new = folder / "eski.xlsx", folder / "yeni.xlsx"
            write_book(old, spec["eski"])
            write_book(new, spec["yeni"])
            for sensitive in (False, True):
                facade = CustomerTrackingFacade(CustomerListReader(), CustomerComparisonExcelExporter(), NoImage(), Drivers(spec["plasiyerler"]), OpenpyxlVisibleTableExporter())
                result = facade.load_and_compare(old, new, case_sensitive=sensitive)
                full = folder / "tam.xlsx"
                facade.export_excel(full, overwrite=True)
                visible = None
                if result.added:
                    facade.export_visible(tuple(reversed(result.added)), "added", folder / "gorunen.xlsx", overwrite=True)
                    visible = cells(folder / "gorunen.xlsx")
                cases.append({**spec, "harfDuyarli": sensitive, "okunanEski": list(CustomerListReader().read(old).customers), "okunanYeni": list(CustomerListReader().read(new).customers), "sonuc": {"eksikler": list(result.missing), "yeniler": list(result.added), "eskiSayisi": result.old_count, "yeniSayisi": result.new_count, "depo": result.depo_name, "mesaj": result.status_text, "dosyaAdi": result.default_output_name}, "tamExcel": cells(full), "gorunenYeniExcel": visible})
        errors = []
        for name, values in [("baslik-yok", [["Yapay"]]), ("gec-baslik", rows(["Yapay"], header=16)), ("kucuk-harf-baslik", [["cari ünvan"], ["Yapay"]]), ("uzun-hucre", rows(["Y" * 513]))]:
            write_book(folder / "hata.xlsx", values)
            try:
                CustomerListReader().read(folder / "hata.xlsx")
            except CustomerDataError as error:
                errors.append({"ad": name, "satirlar": values, "hataTuru": type(error).__name__})
            else:
                raise AssertionError(name)
    files = ["domain/musteri_takip/comparison.py", "domain/musteri_takip/naming.py", "application/customer_tracking/facade.py", "infrastructure/excel/customer_list_reader.py", "infrastructure/export/customer_comparison_exporter.py", "infrastructure/export/visible_table_exporter.py"]
    output = {"kaynak": "BUP_Yonetim — yalnızca yapay veri", "kaynakSha256": {name: hashlib.sha256((root / name).read_bytes()).hexdigest() for name in files}, "senaryolar": cases, "hatalar": errors}
    Path(sys.argv[2]).write_text(json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    if len(sys.argv) > 3:
        katlama = {chr(i): chr(i).casefold() for i in range(0x110000) if chr(i).casefold() != chr(i).lower()}
        Path(sys.argv[3]).write_text(json.dumps(katlama, ensure_ascii=True, indent=2) + "\n", encoding="utf-8")
    print(f"{len(cases)} karşılaştırma ve {len(errors)} hata başvurusu üretildi.")


if __name__ == "__main__":
    main()
