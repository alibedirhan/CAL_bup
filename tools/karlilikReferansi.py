"""Özgün masaüstü okuyucu/facade/çıktı hattıyla yapay kârlılık başvurusu."""
import base64
import hashlib
import json
import sys
import tempfile
from dataclasses import asdict
from datetime import datetime
from pathlib import Path


def main():
    root, output = Path(sys.argv[1]).resolve(), Path(sys.argv[2])
    sys.path.insert(0, str(root))
    from openpyxl import Workbook, load_workbook
    from application.profitability_analysis.facade import ProfitabilityAnalysisFacade
    from application.profitability_analysis.periods import ProfitabilityPeriodService
    from application.profitability_analysis.dashboard import dashboard_statistics, profit_distribution, top_products_by_profit, products_in_category, CATEGORY_LABELS
    from infrastructure.excel.profitability_workbook_reader import OpenpyxlProfitabilityWorkbookReader
    from infrastructure.export.profitability_exporter import OpenpyxlProfitabilityExporter
    from infrastructure.export.profitability_scenario_exporter import OpenpyxlProfitabilityScenarioExporter
    from infrastructure.export.visible_table_exporter import OpenpyxlVisibleTableExporter
    from infrastructure.persistence.profitability_matches import ProfitabilityMatchStore
    from infrastructure.persistence.profitability_periods import PeriodStore

    def workbook(path, rows):
        book = Workbook()
        for row in rows: book.active.append(row)
        for row in book.active:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    cell.data_type = 's'  # Yapay saldırı metni: hesaplanmamış formül değil.
        book.save(path);book.close()
        return base64.b64encode(path.read_bytes()).decode()

    def cells(path):
        book = load_workbook(path, data_only=False)
        result = [{'ad': s.title, 'satirlar': [list(r) for r in s.values]} for s in book]
        book.close();return result

    def dashboard(rows):
        return {'istatistik':asdict(dashboard_statistics(rows)), 'dagilim':asdict(profit_distribution(rows)),
                'ilkUrunler':[asdict(v) for v in top_products_by_profit(rows)],
                'kategoriler':{k:[asdict(v) for v in products_in_category(rows,k)] for k in CATEGORY_LABELS}}

    cases = []
    with tempfile.TemporaryDirectory(prefix='cal-karlilik-reference-') as temp:
        folder = Path(temp)
        reader = OpenpyxlProfitabilityWorkbookReader()
        price = [['Yapay şube alış raporu'], ['Stok İsim','Tarih','Depo','Liste Fiyatı','Fiyat'],
                 [None,None,'YAPAY ALFA',999,'12,345 TL'], [None,None,'YAPAY ALFA',999,99],
                 [None,None,'YAPAY BETA',999,'25,00'], [None,None,'YAPAY İĞÜŞÇ',999,20],
                 [None,None,'YAPAY SIFIR',999,0], [None,None,'YAPAY DEPO',999,1],
                 ['YAPAY HAREKET','01.10.2026','YAPAY ALFA',999,200],
                 [None,None,'=YAPAY()',999,10], [None,None,'YAPAY +',999,'1,234.56']]
        sales = [['Yapay stok dağılımlı rapor'], ['Stok Kodu','Stok İsmi','Satış\nMiktar','Ort.Satış\nFiyat','Satış Tutar'],
                 ['K001','Yapay Alfa',10,'15,00',150],['K002','Yapay Beta',3,20,60],
                 ['K003','Yapay İğüşç',2,20,40],['K004','Yapay Eksik',4,9,36],
                 ['K005','Yapay Takma',1,16,16],['K006','Yapay Alfa',-2,15,-30],
                 ['K007','=YAPAY()',1,12,12],['K008','Yapay +','1.234,50','1,500.25','₺ 1.852.058,625'],
                 [None,'GENEL TOPLAM',999,999,999],[None,None,20,20,400]]
        variants = [
            ('temel',price,sales),
            ('kod-yedegi',price,[['Stok Kodu','Satış Miktar','Ort Satış Fiyat'],['YAPAY ALFA',1,16]]),
            ('baslik-besinci',[[None]]*4+price[1:],[[None]]*4+sales[1:]),
            ('gecersiz-sayi',price,[sales[1],['K1','Yapay Alfa','hatalı','hatalı',None],['K2','Yapay Beta',True,25,True]]),
            ('yalniz-zarar',price,[sales[1],['K1','Yapay Alfa',10,1,10],['K2','Yapay Beta',10,2,20]]),
            ('sifir',price,[sales[1],['K1','Yapay Alfa',0,0,0],['K2','Yapay Beta',1,25,25]]),
            ('yinelenen-isim',price,[sales[1],['K1','Yapay Alfa',1,15,15],['K2','Yapay Alfa',1,15,15]]),
            ('fiyat-kucuk-harf',[price[1],[None,None,'Yapay Alfa',None,10]],[sales[1],['K1','Yapay Alfa',1,15,15]]),
            ('cok-sayfa',price,[sales[1]]+[[f'K{i}',f'Yapay Satır {i:03}',i,15,i*15] for i in range(65)]),
        ]
        for index,(name,p_rows,s_rows) in enumerate(variants):
            p=folder/'fiyat.xlsx';s=folder/'satis.xlsx'
            inputs={'fiyat':workbook(p,p_rows),'satis':workbook(s,s_rows)}
            matches=ProfitabilityMatchStore(folder/f'matches{index}.json')
            facade=ProfitabilityAnalysisFacade(reader,OpenpyxlProfitabilityExporter(),OpenpyxlVisibleTableExporter(),OpenpyxlProfitabilityScenarioExporter(),matches)
            summary=facade.analyze(s,p)
            facade.export(folder/'tam.xlsx',overwrite=True)
            selected=tuple(reversed(summary.rows[::2]))
            indices=[summary.rows.index(row) for row in selected]
            facade.export_visible(selected,folder/'gorunen.xlsx',overwrite=True)
            scenarios=[]
            for values in [(0,0,0),(10,5,-5),(-100,-100,-100),(500,500,500),(7.5,-12.5,33.3)]:
                sc=facade.run_scenario(cost_change_pct=values[0],price_change_pct=values[1],quantity_change_pct=values[2])
                facade.export_scenario(folder/'senaryo.xlsx',overwrite=True)
                scenarios.append({'oranlar':asdict(sc.assumptions),'sonuc':asdict(sc),'excel':cells(folder/'senaryo.xlsx')})
            cases.append({'ad':name,'girdiler':inputs,'fiyatSatirlari':[asdict(r) for r in reader.read_price_report(p)],
                          'satisSatirlari':[asdict(r) for r in reader.read_profitability_report(s)],'ozet':asdict(summary),'genelBakis':dashboard(summary.rows),
                          'eslesme':asdict(facade.match_quality),'tamExcel':cells(folder/'tam.xlsx'), 'gorunen':indices,'gorunenExcel':cells(folder/'gorunen.xlsx'),'senaryolar':scenarios})
            if index==0:
                steps=[]
                facade.propose_match('YAPAY TAKMA','YAPAY ALFA')
                steps.append({'eylem':'oner','ozet':asdict(facade.summary),'eslesme':asdict(facade.match_quality)})
                facade.approve_match('YAPAY TAKMA')
                steps.append({'eylem':'onayla','ozet':asdict(facade.summary),'eslesme':asdict(facade.match_quality)})
                facade.remove_match('YAPAY TAKMA')
                steps.append({'eylem':'kaldir','ozet':asdict(facade.summary),'eslesme':asdict(facade.match_quality)})
                facade.restore_match_change()
                steps.append({'eylem':'geri-al','ozet':asdict(facade.summary),'eslesme':asdict(facade.match_quality)})
                cases[-1]['eslesmeAdimlari']=steps
        errors=[]
        for kind,rows in [('baslik', [['Yapay başlıksız']]),('sutun',[['Stok İsmi','Satış Miktar'],['Yapay Alfa',1]]),('bos',[sales[1]]),('altinci',[[None]]*5+sales[1:])]:
            path=folder/'hata.xlsx';data=workbook(path,rows)
            try: reader.read_profitability_report(path)
            except Exception as e: errors.append({'ad':kind,'bayt':data,'hata':type(e).__name__,'mesaj':str(e)})
        periods=ProfitabilityPeriodService(PeriodStore(folder/'periods.json'),clock=lambda:datetime(2026,10,4,12))
        # İki ayrı saat: kaynağın dönem kimliği ve saat davranışı korunur.
        from application.profitability_analysis.dto import ProfitabilitySummary
        from domain.karlilik.models import ProfitabilityResultRow
        def summary(raw):return ProfitabilitySummary(**{**raw,'rows':tuple(ProfitabilityResultRow(**r) for r in raw['rows']),'unmatched':tuple(raw['unmatched'])})
        first=periods.save_summary('Yapay İlk',summary(cases[0]['ozet']))
        periods._clock=lambda:datetime(2026,10,4,12,0,1)
        second=periods.save_summary('Yapay İkinci',summary(cases[4]['ozet']))
        comparison=periods.compare(first.id,second.id)
        cmp=asdict(comparison)
        cmp['metrics']=[{**asdict(m),'delta':m.delta,'pct_change':m.pct_change} for m in comparison.metrics]
        for key in ('top_gainers','top_losers'):cmp[key]=[{**asdict(m),'delta':m.delta} for m in getattr(comparison,key)]
    files=['domain/karlilik/'+n+'.py' for n in ('__init__','models','analysis','dashboard','matching','period','scenario')]
    files+=['application/profitability_analysis/'+n+'.py' for n in ('dto','facade','ports','dashboard','match_quality','periods','scenario')]
    files+=['infrastructure/excel/profitability_workbook_reader.py','infrastructure/export/profitability_exporter.py','infrastructure/export/profitability_scenario_exporter.py','infrastructure/persistence/profitability_matches.py','infrastructure/persistence/profitability_periods.py']
    output.write_text(json.dumps({'kaynak':'BUP_Yonetim — yalnız yapay Excel','kaynakSha256':{f:hashlib.sha256((root/f).read_bytes()).hexdigest() for f in files},'senaryolar':cases,'hatalar':errors,'donemler':[asdict(first),asdict(second)],'karsilastirma':cmp},ensure_ascii=False,indent=2)+'\n')
    print(f'{len(cases)} okuyucu/analiz, {sum(len(c["senaryolar"]) for c in cases)} senaryo/Excel, {len(errors)} hata başvurusu.')


if __name__=='__main__':main()
