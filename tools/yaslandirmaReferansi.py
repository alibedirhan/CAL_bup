"""Özgün masaüstü okuyucu/domain/facade/çıktı/atama hattıyla yapay yaşlandırma başvurusu.

Kullanım: <kaynak>/.venv/bin/python tools/yaslandirmaReferansi.py <kaynak> <çıktı-json>
Yalnız bu betiğin ürettiği yapay Excel dosyaları okunur; gerçek dosya, config veya
masaüstü atama kaydı kullanılmaz. Kaynak projeye yazılmaz (PYTHONDONTWRITEBYTECODE).
"""
import base64
import hashlib
import json
import re
import sys
import tempfile
import zipfile
from dataclasses import asdict
from datetime import datetime
from pathlib import Path

sys.dont_write_bytecode = True

FILTRE = (
    'Referans Tarihi : 01.10.2026\nYaşlandırma Şekli : Fatura Tarihi Bazlı\n'
    'Cari Kategori 3 : 12 Seçili\nMinimum Bakiye : 100\nBirim Zaman Aralığı : 7 Gün'
)
KOVALAR = ['Açık Hesap', '0-7 Gün', '8-14 Gün', '15-21 Gün', '22-28 Gün', '29-35 Gün', '36-42 Gün',
           '43-49 Gün', '50-56 Gün', '57-63 Gün', '64-70 Gün', '71-77 Gün', '77+ Gün']


def main():
    root, output = Path(sys.argv[1]).resolve(), Path(sys.argv[2])
    sys.path.insert(0, str(root))
    from openpyxl import Workbook, load_workbook
    from application.aging_analysis import AgingAnalysisFacade, VehicleAssignmentService
    from application.aging_analysis.reports import balance_ranking, bucket_totals, vehicle_details
    from infrastructure.excel.aging_workbook_reader import OpenpyxlAgingWorkbookReader
    from infrastructure.export.aging_exporter import OpenpyxlAgingExporter
    from infrastructure.export.visible_table_exporter import OpenpyxlVisibleTableExporter
    from infrastructure.persistence.vehicle_assignments import VehicleAssignmentStore

    def workbook(path, rows, nodim=False):
        book = Workbook()
        for row in rows:
            book.active.append(row)
        for row in book.active:
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith('='):
                    cell.data_type = 's'  # Yapay saldırı metni; hesaplanan formül değil.
        book.save(path)
        book.close()
        if nodim:  # Gerçek muhasebe dışa aktarımları <dimension> taşımayabilir.
            data = {}
            with zipfile.ZipFile(path) as z:
                for item in z.infolist():
                    data[item.filename] = z.read(item.filename)
            with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
                for name, raw in data.items():
                    if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                        raw = re.sub(rb'<dimension[^>]*/>', b'', raw)
                    z.writestr(name, raw)
        return base64.b64encode(path.read_bytes()).decode()

    def cells(path):
        book = load_workbook(path, data_only=False)
        result = [{'ad': s.title, 'satirlar': [list(r) for r in s.values]} for s in book]
        book.close()
        return result

    baslik = ['Cari Kategori 3', 'Cari Ünvan', *KOVALAR]

    def satir(kategori, cari, *degerler):
        return [kategori, cari, *degerler, *[None] * (len(KOVALAR) - len(degerler))]

    temel = [
        ['C01Y.Cari Yaşlandırma Raporu'], [FILTRE], baslik,
        satir('[İZMİR ARAÇ 01]', 'Yapay Alfa Market', 100, '1.234,56', '(50,00)', 0, 0, 300, 0, 0, 0, 0, 0, 0, 0),
        satir('[İZMİR ARAÇ 01]', 'Yapay Beta Gıda', '0', '200', None, '', '-', 'nan', 0, 0, 0, 120, 0, 0, '5,5'),
        satir('İZMİR ARAÇ 2', 'Yapay Gama', '1,234.56', '₺ 100', 'TL 25,75', 0, 0, 0, 0, 0, 0, 0, 0, 0, 0),
        satir('ARAÇ 05', 'Yapay Delta', '-12,5', 0, 0, 0, 0, 0, 0, 0, 0, 0, '1.000.000', 0, 0),
        satir('05', 'Yapay Epsilon', 1e16, 'abc', True, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0),
        satir('[İZMİR ARAÇ 10]', 'Yapay Zeta', 10, 20, 30, 40, 50, 60, 70, 80, 90, 100, 110, 120, 130),
        satir('[İZMİR ARAÇ 10]', None, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1, 1),
        satir('İZMİR ŞUBE DEPO', 'Yapay Depo', 999, 999, 999, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0),
        satir('MERKEZ', 'Yapay Merkez', 999, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0),
        satir('150', 'Yapay Büyük Kod', 999, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0),
        satir(None, 'Yapay Kategorisiz', 999, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0),
        satir('7', '=YAPAY()', 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 77),
        [None] * len(baslik),
    ]
    altin_baslik = ['Cari Kategori 3', 'Cari Ünvan', 'Açık Hesap', '0-7 Gün', '8-14 Gün', '71-77 Gün', '77+ Gün']
    altin = [altin_baslik,
             ['[İZMİR ARAÇ 06]', 'Müşteri A', '100,00', '200', '300,00', 400, '500'],
             ['06', 'Müşteri B', '(50,00)', '', '10,50', None, '5'],
             ['ARAÇ 11', 'Müşteri C', '', '1.234,56', '0', 0, '0'],
             ['İZMİR ŞUBE DEPO', 'Depo Satırı', '999', '999', '999', '999', '999']]
    toplamli = [['Araç', 'Müşteri Adı', 'Açık Hesap', '0-7 Gün', 'Diğer Bakiye', 'Toplam', 'Genel Toplam'],
                ['3', 'Yapay Toplamlı', 10, 20, 30, 60, 60], ['ARAÇ 3', 'Yapay İkinci', 1, 2, 3, 6, 6]]
    tuhaf = [['Kategori 3', 'Firma', '=HYPERLINK("yapay") bakiye', 'Gecikmiş Tutar', '0-7'],
             ['4 ARAÇ', 'Yapay Firma', 5, 6, 7], ['4', 'Yapay Firma 2', '1.000', '2.000,5', '3,333']]
    cok = [baslik] + [satir(f'[İZMİR ARAÇ {n:02d}]', f'Yapay Müşteri {n}-{k}', (n * 37 + k * 11) % 500,
                            (n * k) % 300, 0, 0, 0, (n + k) % 90, 0, 0, 0, (n * 3) % 40, 0, 0, k)
                      for n in range(1, 41) for k in range(1, 4)]
    depolu = [baslik, satir('İZMİR ŞUBE DEPO', 'Yapay Depo', 1), satir('KESİMHANE', 'Yapay Kesim', 2)]
    variants = [
        ('temel', temel, False), ('altin', altin, False), ('boyutsuz', altin, True),
        ('baslik-besinci', [[None]] * 4 + [altin_baslik] + altin[1:], False),
        ('toplam-sutunlari', toplamli, False), ('tuhaf-basliklar', tuhaf, False),
        ('kirk-arac', cok, False), ('arac-yok', depolu, False),
    ]
    cases, errors = [], []
    with tempfile.TemporaryDirectory(prefix='cal-yaslandirma-reference-') as temp:
        folder = Path(temp)
        reader = OpenpyxlAgingWorkbookReader()
        for name, rows, nodim in variants:
            path = folder / 'yaslandirma.xlsx'
            data = workbook(path, rows, nodim)
            report = reader.read_aging_report(path)
            facade = AgingAnalysisFacade(reader, OpenpyxlAgingExporter(), OpenpyxlVisibleTableExporter())
            summary = facade.analyze(path)
            facade.export(folder / 'tam.xlsx', overwrite=True)
            selected = tuple(sorted(summary.vehicles, key=lambda v: (-v.toplam_bakiye, v.arac_no))[::2])
            visible = None
            if selected:
                facade.export_visible(selected, folder / 'gorunen.xlsx', overwrite=True)
                visible = cells(folder / 'gorunen.xlsx')
            cases.append({
                'ad': name, 'bayt': data,
                'okunan': {'satirlar': [asdict(r) for r in report.rows], 'kovalar': list(report.bucket_columns)},
                'ozet': asdict(summary), 'referans': facade.analysis.to_reference(),
                'raporlar': {'detaylar': [asdict(d) for d in vehicle_details(summary.vehicles)],
                             'siralama': [asdict(d) for d in balance_ranking(summary.vehicles)],
                             'kovalar': [list(p) for p in bucket_totals(summary.vehicles)]},
                'tamExcel': cells(folder / 'tam.xlsx'),
                'gorunen': [v.arac_no for v in selected], 'gorunenExcel': visible,
            })
        hata_girdileri = [
            ('baslik-yok', [['Yapay başlıksız'], ['ABC', 100]]),
            ('yalniz-filtre', [[FILTRE], ['ABC', '100']]),
            ('baslik-altinci', [[None]] * 5 + altin),
            ('cari-yok', [['Cari Kategori 3', 'Açık Hesap'], ['ARAÇ 1', 5]]),
            ('kova-yok', [['Cari Kategori 3', 'Cari Ünvan', 'Not'], ['ARAÇ 1', 'Yapay', 'x']]),
            ('veri-yok', [baslik]),
            ('uzun-metin', [baslik, satir('ARAÇ 1', 'Y' * 513, 1)]),
        ]
        for name, rows in hata_girdileri:
            path = folder / 'hata.xlsx'
            data = workbook(path, rows)
            try:
                reader.read_aging_report(path)
                raise AssertionError(name + ' hata vermedi')
            except AssertionError:
                raise
            except Exception as e:  # noqa: BLE001 - kaynak tür ve mesajı kaydedilir
                errors.append({'ad': name, 'bayt': data, 'hata': type(e).__name__, 'mesaj': str(e)})
        bozuk = folder / 'bozuk.xlsx'
        bozuk.write_bytes(b'PK\x03\x04yapay bozuk arsiv')
        for name, raw in [('zip-bozuk', bozuk.read_bytes()), ('excel-degil', b'yapay metin dosyasi')]:
            bozuk.write_bytes(raw)
            try:
                reader.read_aging_report(bozuk)
            except Exception as e:  # noqa: BLE001
                errors.append({'ad': name, 'bayt': base64.b64encode(raw).decode(), 'hata': type(e).__name__, 'mesaj': str(e)})

        store_path = folder / 'atama' / 'atamalar.json'
        store = VehicleAssignmentStore(store_path)
        service = VehicleAssignmentService(store, clock=lambda: datetime(2026, 10, 6, 9, 30))

        def durum(eylem, hata=None):
            def oku(p):
                return p.read_text(encoding='utf-8') if p.is_file() else None
            return {'eylem': eylem, 'hata': hata, 'liste': [asdict(a) for a in service.list_assignments()],
                    'isYuku': service.workload(), 'geriAlinabilir': service.can_restore_last_change(),
                    'kayit': {'guncel': oku(store_path), 'yedek': oku(store.backup_path)}}

        steps = [durum('bos')]
        plans = [
            ('ata', {'arac_no': ' 5 ', 'sorumlu': ' Yapay Sorumlu A ', 'email': 'a@yapay.test', 'telefon': '0500 000 00 00', 'departman': 'Yapay Satış', 'notlar': 'Yapay not'}),
            ('ata', {'arac_no': '7', 'sorumlu': 'Yapay Sorumlu B'}),
            ('ata', {'arac_no': '10', 'sorumlu': 'Yapay Sorumlu A'}),
            ('ata', {'arac_no': '5', 'sorumlu': 'Yapay Sorumlu C', 'notlar': '=YAPAY()'}),
            ('ata', {'arac_no': '9', 'sorumlu': '   '}),
            ('kaldir', {'arac_no': '7'}),
            ('geri-al', {}),
            ('geri-al', {}),
            ('kaldir', {'arac_no': '99'}),
        ]
        for eylem, alanlar in plans:
            hata = None
            try:
                if eylem == 'ata':
                    service.assign(**alanlar)
                elif eylem == 'kaldir':
                    service.remove(alanlar['arac_no'])
                else:
                    hata = None if service.restore_last_change() else 'geri-alinamadi'
            except Exception as e:  # noqa: BLE001
                hata = type(e).__name__ + ': ' + str(e)
            steps.append({**durum(eylem, hata), 'alanlar': alanlar})

    files = ['domain/yaslandirma/' + n + '.py' for n in ('__init__', 'models', 'analysis', 'buckets', 'numbers', 'vehicles', 'reports', 'assignments')]
    files += ['application/aging_analysis/' + n + '.py' for n in ('__init__', 'dto', 'facade', 'ports', 'reports', 'assignments')]
    files += ['infrastructure/excel/aging_workbook_reader.py', 'infrastructure/export/aging_exporter.py',
              'infrastructure/persistence/vehicle_assignments.py', 'core/versioned_json_store.py',
              'core/visible_export.py', 'application/visible_export.py', 'core/cancellation.py',
              'infrastructure/export/visible_table_exporter.py', 'infrastructure/export/atomic_publish.py']
    output.write_text(json.dumps({
        'kaynak': 'BUP_Yonetim — yalnız yapay Excel ve geçici atama kaydı',
        'kaynakSha256': {f: hashlib.sha256((root / f).read_bytes()).hexdigest() for f in files},
        'senaryolar': cases, 'hatalar': errors, 'atamaAdimlari': steps,
    }, ensure_ascii=False, indent=2) + '\n')
    print(f'{len(cases)} okuma/analiz/Excel, {len(errors)} hata, {len(steps)} atama adımı.')


if __name__ == '__main__':
    main()
