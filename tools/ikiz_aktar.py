# -*- coding: utf-8 -*-
"""Eski aracın Python ikizini gerçek dosyalarla çalıştırıp sonucu JSON'a yazar.

TypeScript kurallarının ikizle birebir aynı sonucu verdiğini kanıtlamak için
tests/altin/ikiz.test.ts bu dosyayı okur. Çıktı ornekler/ altına yazılır ve
git'e girmez (şirket verisi).

Çalıştırma (eski aracın sanal ortamıyla):
    ../bupilic-rapor-araci/.venv/bin/python tools/ikiz_aktar.py [GUN]
GUN verilmezse depo kontrol dosyasındaki son gün sayfası kullanılır.
"""
import glob, json, os, sys, unicodedata
from datetime import date, datetime

KOK = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ESKI = os.path.join(os.path.dirname(KOK), 'bupilic-rapor-araci')
ORN = os.path.join(KOK, 'ornekler')
sys.path.insert(0, ESKI)
sys.path.insert(0, os.path.join(ESKI, 'src'))

import openpyxl  # noqa: E402
from ikiz.led import oku_d01, oku_sayim, oku_sube_alis, sayfa_tarihi_yilli  # noqa: E402
from ikiz.depo_kontrol import hesapla, liste_oku  # noqa: E402

BUGUN = date(2026, 10, 1)


def bul(*desenler):
    for f in sorted(glob.glob(os.path.join(ORN, '*'))):
        ad = unicodedata.normalize('NFC', os.path.basename(f)).upper()
        if f.lower().endswith(('.xlsx', '.xlsm')) and all(d.upper() in ad for d in desenler):
            return f
    raise SystemExit(f'ornekler/ içinde {desenler} bulunamadı')


def deger(v):
    if isinstance(v, (datetime, date)):
        return {'$tarih': v.isoformat()}
    return v


def izgara(yol):
    wb = openpyxl.load_workbook(yol, data_only=True)
    olus = wb.properties.created
    return {
        'dosyaAdi': unicodedata.normalize('NFC', os.path.basename(yol)),
        'olusturulma': olus.isoformat() if isinstance(olus, datetime) else None,
        'sayfalar': [
            {'ad': ws.title, 'satirlar': [[deger(c) for c in r] for r in ws.iter_rows(values_only=True)]}
            for ws in wb.worksheets
        ],
    }


def iso(d):
    return d.isoformat() if d else None


def kaynak(kv):
    return {'tarih': iso(kv.tarih), 'baslangicTarihi': iso(kv.baslangic_tarihi), 'dipToplam': kv.dip_toplam,
            'dipToplamVar': kv.dip_toplam_var, 'toplam': kv.toplam(), 'urunSayisi': len(kv.miktarlar)}


def main():
    yollar = dict(d01=bul('D01'), sayim=bul('SAYIM'), sube=bul('ALI'), hedef=bul('DEPO', 'KONTROL'))
    hedef = openpyxl.load_workbook(yollar['hedef'], data_only=True)
    gun = sys.argv[1] if len(sys.argv) > 1 else hedef.sheetnames[-1]
    onceki = hedef.worksheets[hedef.sheetnames.index(gun) - 1]

    d01 = oku_d01(yollar['d01'])
    sayim = oku_sayim(yollar['sayim'], bugun=BUGUN)
    sube = oku_sube_alis(yollar['sube'])
    liste = liste_oku(onceki)
    s = hesapla(liste, d01, sayim, sube)

    cikti = {
        'bugun': BUGUN.isoformat(),
        'gun': gun,
        'oncekiGun': onceki.title,
        'yeniTarih': iso(sayfa_tarihi_yilli(gun, BUGUN)),
        'oncekiTarih': iso(sayfa_tarihi_yilli(onceki.title, BUGUN)),
        'dosyalar': {k: izgara(yollar[k]) for k in ('d01', 'sayim', 'sube')},
        'liste': liste,
        'ikiz': {
            'kaynaklar': {'d01': kaynak(d01), 'sayim': kaynak(sayim), 'sube': kaynak(sube)},
            'satirlar': [[x.ad, x.b, x.d, x.eklendi] for x in s.satirlar],
            'eklenenler': s.eklenenler,
            'sifirEksikler': s.sifir_eksikler,
            'tekrarlar': [[i + 4, j + 4, ad] for i, j, ad in s.tekrarlar],
            'donukToplam': s.donuk_toplam,
            'bToplam': s.b_toplam,
            'dToplam': s.d_toplam,
            'kontroller': [list(k) for k in s.kontroller],
        },
    }
    yol = os.path.join(ORN, 'ikiz.json')
    with open(yol, 'w', encoding='utf-8') as f:
        json.dump(cikti, f, ensure_ascii=False)
    print(f'{yol}: {gun} ({onceki.title} sayfasından), {len(s.satirlar)} satır, '
          f'{len(s.eklenenler)} eklenen, kontroller: {[k[3] for k in s.kontroller]}')


if __name__ == '__main__':
    main()
